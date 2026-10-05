# -*- coding: utf-8 -*-
"""
Menyusun naskah INOVATIF (Bahasa Indonesia & Inggris) langsung dari tabel/gambar hasil eksperimen.
  python src/build_manuscript.py --template "project4/Template Inovatif 2022 ok.docx" --lang id|en|both
Template .doc dikonversi ke .docx terlebih dahulu (LibreOffice: soffice --headless --convert-to docx).
"""
import os, re, copy, json, argparse
import pandas as pd, numpy as np
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from refs import REFS

ap = argparse.ArgumentParser()
HERE = os.path.dirname(os.path.abspath(__file__))
ap.add_argument("--root", default=os.path.abspath(os.path.join(HERE, "..")), help="folder proyek yang berisi results/ (mis. .../project4)")
ap.add_argument("--template", default=None)
ap.add_argument("--lang", default="both")
ap.add_argument("--out_dir", default=None)
args, _ = ap.parse_known_args()
ROOT = os.path.abspath(args.root)
args.template = args.template or os.path.join(ROOT, "Template Inovatif 2022 ok.docx")
args.out_dir = args.out_dir or os.path.join(ROOT, "manuscript")
os.makedirs(args.out_dir, exist_ok=True)
TAB = os.path.join(ROOT, "results", "tables")
rd = lambda n, **k: pd.read_csv(os.path.join(TAB, n + ".csv"), **k)

# ------------------------------------------------------------------ angka hasil (semua dari CSV)
N = {}
S = json.load(open(os.path.join(ROOT, "results", "run_summary.json")))
N["S"] = S
agg = rd("S06_perbandingan_metode_numerik", header=[0, 1], index_col=0)
N["agg"] = agg
N["T07"] = rd("T07_perbandingan_metode"); N["T08"] = rd("T08_recall_per_tipe_anomali").set_index("tipe")
N["T09"] = rd("T09_uji_statistik"); N["T10"] = rd("T10_sensitivitas")
N["T11"] = rd("T11_ablasi_fitur", header=[0, 1], index_col=0)
N["T12"] = rd("T12_hasil_data_riil_per_tahun"); N["T03"] = rd("T03_prevalensi_aturan_per_tahun")
N["T04"] = rd("T04_skor_kualitas_data"); N["DRIFT"] = rd("T03b_pergeseran_kode"); N["T01"] = rd("T01_inventaris_data")
N["cross"] = rd("S07_irisan_aturan_ML").set_index("aturan"); N["T13"] = rd("T13_top20_anomali_hibrida"); N["T13b"] = rd("T13b_top20_anomali_ML_saja")
N["S08"] = rd("S08_alasan_utama_ML_saja"); N["TM"] = rd("S05_waktu_komputasi").mean(numeric_only=True)
N["T02"] = rd("T02_katalog_aturan"); N["T06"] = rd("T06_tipe_injeksi_anomali")
def m(meth, met, sd=False):
    v = agg.loc[meth, (met, "mean")]; s = agg.loc[meth, (met, "std")]
    return f"{v:.3f} ± {s:.3f}" if sd else f"{v:.3f}"
N["m"] = m

# ------------------------------------------------------------------ format angka per bahasa
def fmt_num(x, lang, d=0):
    s = f"{x:,.{d}f}"
    if lang == "id": s = s.replace(",", "§").replace(".", ",").replace("§", ".")
    return s
def fmt_dec(x, lang, d=3):
    s = f"{x:.{d}f}"
    return s.replace(".", ",") if lang == "id" else s

# ------------------------------------------------------------------ mesin dokumen
class Builder:
    def __init__(self, template, lang):
        self.doc = Document(template); self.lang = lang
        self.fig_n = 0; self.tab_n = 0; self.order = []
        body = self.doc.element.body; els = list(body.iterchildren())
        for el in els[15:]:
            if el.tag != qn("w:sectPr"): body.remove(el)
    # --- teks kaya: **tebal**, *miring*, ^sup^, [@kunci; @kunci]
    def cite(self, txt):
        def rep(mm):
            keys = [k.strip().lstrip("@") for k in mm.group(1).split(";")]
            nums = []
            for k in keys:
                assert k in REFS, k
                if k not in self.order: self.order.append(k)
                nums.append(self.order.index(k) + 1)
            nums = sorted(set(nums)); out = []; i = 0
            while i < len(nums):
                j = i
                while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1: j += 1
                out.append(f"{nums[i]}–{nums[j]}" if j - i >= 2 else ",".join(str(x) for x in nums[i:j + 1]))
                i = j + 1
            return "[" + ",".join(out) + "]"
        return re.sub(r"\[(@[^\]]+)\]", rep, txt)
    def runs(self, p, txt, size=10, bold=False, italic=False):
        txt = self.cite(txt)
        for tok in re.split(r"(\*\*.+?\*\*|\*[^*]+?\*|\^[^^]+?\^|~[^~]+?~)", txt):
            if not tok: continue
            b, it, sup, sub = bold, italic, False, False
            if tok.startswith("**"): tok, b = tok[2:-2], True
            elif tok.startswith("*"): tok, it = tok[1:-1], True
            elif tok.startswith("^"): tok, sup = tok[1:-1], True
            elif tok.startswith("~"): tok, sub = tok[1:-1], True
            r = p.add_run(tok); f = r.font; f.name = "Times New Roman"; f.size = Pt(size); f.bold = b; f.italic = it
            f.superscript = sup; f.subscript = sub
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        return p
    def fmt(self, p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=0, indent=None, hanging=None, keep=False):
        pf = p.paragraph_format; p.alignment = align; pf.space_before = Pt(before); pf.space_after = Pt(after)
        pf.line_spacing = 1.0
        if indent is not None: pf.first_line_indent = indent
        if hanging is not None: pf.left_indent = hanging; pf.first_line_indent = -hanging
        if keep: pf.keep_with_next = True
        return p
    def para(self, txt, **k):
        size = k.pop("size", 10); bold = k.pop("bold", False); italic = k.pop("italic", False)
        p = self.doc.add_paragraph(); self.fmt(p, **k); self.runs(p, txt, size, bold, italic); return p
    def blank(self): return self.para("")
    def h1(self, txt): self.blank(); return self.para(txt, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, keep=True, after=3)
    def h2(self, txt): return self.para(txt, bold=True, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, keep=True, before=4, after=2)
    def eq(self, txt, n):
        t = self.doc.add_table(rows=1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.columns[0].width = Cm(14); t.columns[1].width = Cm(1.4)
        c0, c1 = t.rows[0].cells; c0.width = Cm(14); c1.width = Cm(1.4)
        self.fmt(c0.paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER, 2, 2); self.runs(c0.paragraphs[0], txt, 10, italic=False)
        self.fmt(c1.paragraphs[0], WD_ALIGN_PARAGRAPH.RIGHT, 2, 2); self.runs(c1.paragraphs[0], f"({n})", 10)
    def figure(self, path, caption, width=14.5):
        self.fig_n += 1
        p = self.doc.add_paragraph(); self.fmt(p, WD_ALIGN_PARAGRAPH.CENTER, 4, 0, keep=True)
        p.add_run().add_picture(path, width=Cm(width))
        lab = "Gambar" if self.lang == "id" else "Figure"
        self.para(f"{lab} {self.fig_n}: {caption}", size=9, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
        return self.fig_n
    def table(self, df, caption, widths=None, size=8, bold_last=False, bold_rows=()):
        self.tab_n += 1
        lab = "Tabel" if self.lang == "id" else "Table"
        self.para(f"{lab} {self.tab_n}: {caption}", size=8, align=WD_ALIGN_PARAGRAPH.CENTER, before=4, keep=True)
        t = self.doc.add_table(rows=len(df) + 1, cols=len(df.columns)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = t._tbl.tblPr; borders = OxmlElement("w:tblBorders")
        for side in ["top", "bottom"]:
            e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "8"); e.set(qn("w:color"), "000000"); borders.append(e)
        for side in ["left", "right", "insideH", "insideV"]:
            e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "nil"); borders.append(e)
        tblPr.append(borders)
        for j, c in enumerate(df.columns):
            cell = t.rows[0].cells[j]; cell.text = ""
            self.fmt(cell.paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER); self.runs(cell.paragraphs[0], str(c), size, bold=True)
            tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement("w:tcBorders"); e = OxmlElement("w:bottom")
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6"); b.append(e); tcPr.append(b)
        for i, row in enumerate(df.itertuples(index=False), start=1):
            for j, v in enumerate(row):
                cell = t.rows[i].cells[j]; cell.text = ""
                al = WD_ALIGN_PARAGRAPH.LEFT if j == 0 or isinstance(v, str) and len(v) > 14 else WD_ALIGN_PARAGRAPH.CENTER
                self.fmt(cell.paragraphs[0], al)
                self.runs(cell.paragraphs[0], str(v), size, bold=(bold_last and i == len(df)) or (i - 1) in bold_rows)
        if widths:
            t.autofit = False
            tot = sum(widths); widths = [w * min(1.0, 16.0 / tot) for w in widths]
            grid = t._tbl.tblGrid
            for j, gc in enumerate(grid.findall(qn("w:gridCol"))): gc.set(qn("w:w"), str(int(widths[j] * 567)))
            for j, w in enumerate(widths):
                for r in t.rows: r.cells[j].width = Cm(w)
            tw = OxmlElement("w:tblW"); tw.set(qn("w:w"), str(int(sum(widths) * 567))); tw.set(qn("w:type"), "dxa")
            old_ = tblPr.find(qn("w:tblW"))
            if old_ is not None: tblPr.remove(old_)
            tblPr.append(tw)
            lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
        for r in t.rows:  # rapatkan sel
            for c in r.cells:
                tcPr = c._tc.get_or_add_tcPr(); mar = OxmlElement("w:tcMar")
                for side in ("top", "bottom"):
                    e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), "15"); e.set(qn("w:type"), "dxa"); mar.append(e)
                tcPr.append(mar)
        self.blank()
        return self.tab_n
    # --- halaman judul
    def set_par(self, idx, segs, size):
        p = self.doc.paragraphs[idx]
        for r in list(p.runs): r._element.getparent().remove(r._element)
        for txt, bold, sup in segs:
            r = p.add_run(txt); r.font.name = "Times New Roman"; r.font.size = Pt(size); r.font.bold = bold; r.font.superscript = sup
    def front(self, title, authors, affs, emails):
        self.set_par(0, [(title, True, False)], 14)
        for i in (2, 3, 6, 10, 13): self.set_par(i, [("", False, False)], 10)
        segs = []
        for k, (nm, sup) in enumerate(authors):
            segs += [(nm, True, False), (sup, True, True)]
            if k < len(authors) - 1: segs.append((", ", True, False))
        self.set_par(5, segs, 12)
        self.set_par(8, [(affs[0][0], False, True), (affs[0][1], False, False)], 10)
        self.set_par(9, [(affs[1][0], False, True), (affs[1][1], False, False)], 10)
        self.set_par(12, [(emails, False, False)], 10)
    def references(self, head):
        self.h1(head)
        for i, k in enumerate(self.order, 1):
            p = self.doc.add_paragraph(); self.fmt(p, hanging=Cm(0.9))
            r = p.add_run(f"[{i}]\t{REFS[k]}"); r.font.name = "Times New Roman"; r.font.size = Pt(10)
    def save(self, path):
        self.doc.save(path); print("saved", path, "| refs:", len(self.order), "| fig:", self.fig_n, "| tab:", self.tab_n)

if __name__ == "__main__":
    import content_id, content_en
    langs = ["id", "en"] if args.lang == "both" else [args.lang]
    for lg in langs:
        b = Builder(args.template, lg)
        (content_id if lg == "id" else content_en).build(b, N, ROOT)
        name = "INOVATIF_Deteksi_Anomali_SAMSAT_PKB_Banten_ID.docx" if lg == "id" else "INOVATIF_SAMSAT_Anomaly_Detection_PKB_Banten_EN.docx"
        b.save(os.path.join(args.out_dir, name))

def build_supplement(lang):
    """Materi suplemen: gambar & tabel tambahan yang tidak dimuat di naskah utama (batas 10 halaman)."""
    d = Document(); st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(10)
    fd = os.path.join(ROOT, "results", "figures" if lang == "id" else "figures_en")
    T = (lambda i, e: i if lang == "id" else e)
    d.add_heading(T("Materi Suplemen – Deteksi Anomali Hibrida pada Data PKB SAMSAT Banten",
                    "Supplementary Material – Hybrid Anomaly Detection in SAMSAT Banten PKB Data"), 1)
    figs = [("F12_ablasi_fitur", T("Ablasi kontribusi fitur berbasis domain terhadap PR-AUC (5 ulangan)", "Ablation of domain-feature contribution to PR-AUC (5 runs)")),
            ("F07_kurva_ROC_PR", T("Kurva ROC dan precision–recall (ulangan pertama)", "ROC and precision–recall curves (first run)")),
            ("F10_distribusi_skor", T("Distribusi skor AE dan IF: normal vs tipe anomali", "AE and IF score distributions: normal vs anomaly types")),
            ("F11_sensitivitas", T("Analisis sensitivitas hiperparameter dan strategi fusi", "Hyper-parameter and fusion-strategy sensitivity")),
            ("F04_skor_kualitas_data", T("Kartu skor kualitas data per dimensi", "Data-quality scorecard by dimension")),
            ("F05_nilai_pkb_per_jenis", T("Nilai PKB per jenis kendaraan dan batas R09", "PKB value per vehicle type and R09 limits")),
            ("F06_matriks_kode_mutasi", T("Matriks kode permohonan × mutasi (setelah harmonisasi)", "Application × mutation code matrix (after harmonisation)")),
            ("F14_proyeksi_PCA", T("Proyeksi PCA ruang fitur dengan kategori temuan", "PCA projection of the feature space by finding category")),
            ("F15_waktu_komputasi", T("Waktu komputasi rata-rata", "Mean computation time"))]
    for i, (f, cap) in enumerate(figs, 1):
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(os.path.join(fd, f + ".png"), width=Cm(15))
        c = d.add_paragraph(f"{T('Gambar', 'Figure')} S{i}: {cap}"); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for k, (name, cap) in enumerate([("T09_uji_statistik", T("Uji Wilcoxon (Hibrida vs pembanding, PR-AUC) dan Friedman", "Wilcoxon (hybrid vs competitors, PR-AUC) and Friedman tests")),
                                     ("T10_sensitivitas", T("Analisis sensitivitas", "Sensitivity analysis")),
                                     ("T03b_pergeseran_kode", T("Pergeseran kode per atribut dan tahun (JSD)", "Code drift per field and year (JSD)")),
                                     ("S01_profil_kode_mohon_mutasi", T("Profil pasangan kode permohonan–mutasi (2021–2023)", "Application–mutation code-pair profile (2021–2023)"))], 1):
        df = rd(name)
        d.add_paragraph(f"{T('Tabel', 'Table')} S{k}: {cap}").alignment = WD_ALIGN_PARAGRAPH.CENTER
        t = d.add_table(rows=len(df) + 1, cols=len(df.columns)); t.style = "Table Grid"
        for j, c in enumerate(df.columns): t.rows[0].cells[j].text = str(c)
        for i, row in enumerate(df.itertuples(index=False), 1):
            for j, v in enumerate(row): t.rows[i].cells[j].text = (f"{v:.4g}" if isinstance(v, float) else str(v))
        for r in t.rows:
            for c in r.cells:
                for p in c.paragraphs:
                    for rr in p.runs: rr.font.size = Pt(7.5)
        d.add_paragraph("")
    out = os.path.join(args.out_dir, "Suplemen_Deteksi_Anomali_SAMSAT_ID.docx" if lang == "id" else "Supplementary_SAMSAT_Anomaly_Detection_EN.docx")
    d.save(out); print("saved", out)

if __name__ == "__main__":
    for lg in (["id", "en"] if args.lang == "both" else [args.lang]):
        build_supplement(lg)
